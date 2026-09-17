"""Exact-source ingestion and governed retrieval; no live model or Pinecone required.

These tests prove grounding and prompt constraints, not stochastic model compliance.
"""
from hashlib import sha256
from io import StringIO
from pathlib import Path
from unittest.mock import patch

import docx
import pytest
from django.core.management import call_command

from apps.ai_engine.services.knowledge_retriever import KnowledgeRetriever
from apps.ai_engine.services.system_prompt_builder import _format_context
from apps.knowledge_core.models import KnowledgeDocument
from apps.knowledge_core.services.document_loader import load_document_text
from apps.knowledge_core.services.ingestion_pipeline import ingest_document
from apps.knowledge_core.services.metadata_tagger import build_chunk_metadata
from apps.knowledge_core.source_manifest import people_sources_for, policy_for

ROOT = Path(__file__).resolve().parents[2]
KANG = 'itriX_Knowledge_Core_Kang_Myungjoo_v1.0.docx'
PARK = 'itriX_Knowledge_Core_Park_Junhu_v1.0.docx'
ASTOP = 'itriX_ASTOP_Comparative_Knowledge_Core_v1.1.docx'
HASHES = {
    KANG: 'abc65cc135a0a50da98843cdb1185c9fd8c318381a88f94ba6af5032f0eafa3c',
    PARK: 'e7fd10d47e8050b5aabeff5204e7d770a1452ac6a7915fce8358d2289eab1c97',
    ASTOP: 'a5c91a732e38e53a2dab57a3d43825b4a18b2b60407e21bff7040bf1ec6c4764',
}


@pytest.fixture
def corpus(db, settings):
    settings.ENABLE_AI_ENGINE = False
    call_command('register_knowledge_docs', stdout=StringIO())
    docs = {}
    for name in (*HASHES, 'itrix_product_canonical_v3_5.md', 'astop_prism_public_safe_v2_3.md'):
        document = KnowledgeDocument.objects.get(file_path__endswith='/' + name)
        result = ingest_document(document)
        assert result.ok, result.error
        docs[name] = document
    return docs


@pytest.mark.parametrize('name', HASHES)
def test_exact_binary_and_paragraph_table_extraction(name):
    path = ROOT / 'knowledge_docs/controlled_public' / name
    assert sha256(path.read_bytes()).hexdigest() == HASHES[name]
    parsed = docx.Document(path)
    loaded = load_document_text(str(path))
    assert len(loaded) > 10000
    assert any(p.text.strip() and p.text.strip() in loaded for p in parsed.paragraphs)
    assert parsed.tables
    for table in parsed.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    assert cell.text.strip() in loaded


@pytest.mark.parametrize('name', HASHES)
def test_registration_and_reingestion_are_idempotent(corpus, name):
    document = corpus[name]
    before = list(document.chunks.values_list('vector_id', flat=True))
    call_command('register_knowledge_docs', stdout=StringIO())
    document.refresh_from_db()
    assert document.ingestion_status == 'COMPLETE'
    assert KnowledgeDocument.objects.filter(file_path=document.file_path).count() == 1
    assert document.is_current
    assert document.disclosure_level == 'controlled_public'
    assert document.namespace == ('astop' if name == ASTOP else 'company')
    assert document.verified_at.date().isoformat() == policy_for(name).verified_date
    assert len(document.canonical_rule) <= 512
    assert ingest_document(document).ok
    assert list(document.chunks.values_list('vector_id', flat=True)) == before
    assert document.content_hash == HASHES[name]
    chunk = document.chunks.first()
    metadata = build_chunk_metadata(document=document, chunk=chunk)
    assert metadata['verified_at'].startswith(policy_for(name).verified_date)
    assert metadata['canonical_entities'] == list(policy_for(name).canonical_entities)
    assert metadata['canonical_rule'] == document.canonical_rule
    assert not document.supersedes
    for other in ('ASTOP_Productization_GTM_Plan_v2.3.docx', 'prism-paper-current_v2.pdf',
                  'astop_prism_public_safe_v2_3.md', 'ASTOP_Technical_Capabilities_Current_v0.3.1.md'):
        assert KnowledgeDocument.objects.get(file_path__endswith='/' + other).is_current


PEOPLE_CASES = [
    ('Who is Myungjoo Kang?', KANG), ('Who is Kang Myungjoo?', KANG),
    ('Who is Professor Kang?', KANG), ('Who is 강명주?', KANG),
    ('Who is the CEO of itriX?', KANG), ('Did Kang invent AXIOM?', KANG),
    ('Did Kang invent CRE?', KANG), ("What is Kang's connection to FQNM?", KANG),
    ('How many papers does Kang have?', KANG), ('Is Kang currently KSIAM president?', KANG),
    ('Who is Park Junhu?', PARK), ('Who is Junhu Park?', PARK), ('Who is 박준후?', PARK),
    ('Who leads itriX R&D?', PARK), ('Who wrote AXIOM?', PARK),
    ('Who is behind FQNM?', PARK), ('What did Park do on AXIOM?', PARK),
    ("What is Park's role in CRE?", PARK), ("What is Park's role in FQNM?", PARK),
    ('Did Park invent all itriX technology?', PARK), ('Is Park a doctor?', PARK),
    ('How many patents does Park have?', PARK), ("What is Park's home address?", PARK),
    ("What is Park's resident registration number?", PARK),
]

@pytest.mark.parametrize('query,name', PEOPLE_CASES)
def test_people_retrieval_and_constraints(corpus, query, name):
    chunks = KnowledgeRetriever().retrieve(query, journey_stage='ARRIVED')
    selected = [c for c in chunks if c['document_id'] == str(corpus[name].id)]
    assert selected, query
    assert all(c['namespace'] == 'company' and c['disclosure_level'] == 'controlled_public' for c in selected)
    assert all(c['verified_at'].startswith('2026-09-08') for c in selected)
    prompt = _format_context(chunks)
    assert corpus[name].canonical_rule in prompt
    for boundary in corpus[name].prohibited_messages:
        assert boundary in prompt


ASTOP_QUERIES = [
    'How is ASTOP different from prompt caching?', 'Is ASTOP just prompt caching?',
    'How is ASTOP different from NVIDIA SoL-Pi?', 'Is ASTOP better than SoL-Pi?',
    'Does ASTOP always save 51.9–84.5%?', 'Is ASTOP 51.9–84.5% cheaper than SoL-Pi?',
    'Can we add ASTOP, SoL-Pi and Anthropic savings?', 'Why not just use webhooks?',
    'Does ASTOP itself consume compute?', 'Will ASTOP improve model intelligence?',
    'When is ASTOP a weak fit?',
]

@pytest.mark.parametrize('query', ASTOP_QUERIES)
def test_astop_comparison_retrieval_and_constraints(corpus, query):
    chunks = KnowledgeRetriever().retrieve(query, journey_stage='ARRIVED')
    selected = [c for c in chunks if c['document_id'] == str(corpus[ASTOP].id)]
    assert selected, query
    assert all(c['namespace'] == 'astop' for c in selected)
    assert all(c['verified_at'].startswith('2026-09-16') for c in selected)
    prompt = _format_context(chunks)
    assert corpus[ASTOP].canonical_rule in prompt
    for boundary in corpus[ASTOP].prohibited_messages:
        assert boundary in prompt


@pytest.mark.parametrize('name', (KANG, PARK))
def test_all_authorized_aliases(corpus, name):
    for alias in policy_for(name).aliases:
        query = f'Who is {alias}?'
        assert name in people_sources_for(query)
        chunks = KnowledgeRetriever().retrieve(query, journey_stage='ARRIVED')
        assert any(c['document_id'] == str(corpus[name].id) for c in chunks), query


def test_person_resolution_never_expands_namespace_or_disclosure(corpus):
    retriever = KnowledgeRetriever()
    assert not retriever.retrieve('Who is 강명주?', namespaces=['general'], journey_stage='ARRIVED')
    document = corpus[KANG]
    document.approved_audience = ['internal']
    document.save()
    chunks = retriever.retrieve('Who is Myungjoo Kang?', journey_stage='ARRIVED')
    assert not any(c['document_id'] == str(document.id) for c in chunks)


def test_vector_alias_retrieval_uses_existing_namespaces_and_db_governance(corpus, settings):
    settings.ENABLE_AI_ENGINE = True
    row = corpus[KANG].chunks.first()
    with patch('apps.ai_engine.services.knowledge_retriever.Embedder') as embedder, patch(
        'apps.ai_engine.services.knowledge_retriever.PineconeQueryClient'
    ) as client:
        embedder.return_value.embed_one.return_value = [0.0]
        client.return_value.query.side_effect = lambda **kw: (
            [{'id': row.vector_id, 'score': .99}] if kw['namespace'] == 'company' else []
        )
        chunks = KnowledgeRetriever().retrieve('Who is 강명주?', journey_stage='ARRIVED')
        assert 'Myungjoo Kang' in embedder.return_value.embed_one.call_args.args[0]
        assert chunks[0]['document_id'] == str(corpus[KANG].id)
        assert chunks[0]['canonical_rule'] == corpus[KANG].canonical_rule
        assert chunks[0]['retrieval_backend'] == 'pinecone'
        assert {c.kwargs['namespace'] for c in client.return_value.query.call_args_list} == {
            'company', 'astop', 'technology', 'alpha-compute', 'alpha-core', 'proofs', 'licensing'
        }
